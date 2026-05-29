# -*- coding: utf-8 -*-
"""Append 'Appendix A - Event Catalogue' (19 events) to the report docx."""
import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "999999")
        borders.append(el)
    tblPr.append(borders)

HERE = os.path.dirname(os.path.abspath(__file__))
REPORT = os.path.normpath(os.path.join(HERE, "..", "f1_telemetry.docx"))
IMG_W = Inches(6.3)

# ---- Event data ----------------------------------------------------------
# (img, name, why, handling)
CAR = [
 ("ev01_state-change", "state-change",
  "Umbrella event emitted whenever the per-car finite state machine accepts a "
  "transition (any from -> to). It is produced inside Car._applyTransition(), so "
  "every other lifecycle change below is accompanied by a state-change.",
  "The simulator also publishes a retained state message on the .../state topic. "
  "The state-change event itself (details {from, to, reason}) is published on the "
  "per-car events topic at QoS 1, stored in the events collection and pushed to "
  "the dashboard, which updates the car colour/status badge."),
 ("ev02_pit-entry", "pit-entry",
  "The car enters the pit lane: conditions.js detects tyre wear >= 0.60 while the "
  "car is inside the pit-entry window (trackPos >= 0.95), firing the tire-service "
  "trigger and the RUNNING -> PIT transition.",
  "A pit-entry event (details {lap}) is published alongside the state-change, "
  "persisted and shown on the dashboard timeline. The car keeps moving through "
  "the pit lane until the pit box is reached."),
 ("ev03_pit-stop", "pit-stop",
  "Fired when the car, already in PIT, crosses the pit-box position (trackPos ~= "
  "0.985) and the tyre-change service is performed. Modern-F1 rule: tyres only, "
  "no refuelling; the car must rejoin on a DIFFERENT compound.",
  "Tyre temperature is reset, a different compound is drawn from a dedicated PRNG "
  "stream and the wear accumulator is zeroed. The pit-stop event carries "
  "{duration (2.0-3.5 s), tyreCompound, service:'tire-change', refuelling:false} "
  "and is persisted/broadcast like every car event."),
 ("ev04_pit-exit", "pit-exit",
  "The pit stop is complete: the pit timer (2.0-3.5 s, sampled per stop) has "
  "elapsed and the pit-exit point has been reached, firing the pit-out trigger "
  "and the PIT -> RUNNING transition.",
  "A pit-exit event (details {lap}) is emitted together with the state-change; "
  "the dashboard returns the car to the RUNNING colour and the lap/sector timers "
  "resume."),
 ("ev05_fault", "fault",
  "The car enters FAULT. Three independent causes: tyre overheat (max tyre temp "
  "> 180 C for 3 consecutive ticks), random engine failure (probability 1e-4 per "
  "tick) or internal-error (an exception thrown inside Car.tick(), which forces "
  "FAULT from any state).",
  "On entering FAULT a 5 s diagnosis timer starts (onEnterFault). The fault event "
  "(details {reason}) is published, persisted and shown on the dashboard. A car "
  "in FAULT is a hazard and feeds the YELLOW / VSC Race Control triggers."),
 ("ev06_retirement", "retirement",
  "Permanent withdrawal (absorbing RETIRED state). Reached when the 5 s FAULT "
  "diagnosis elapses (every fault is modelled as unrecoverable) or when Race "
  "Control issues a manual-retire from any non-terminal state.",
  "The retirement event (details {reason}) is published, persisted and rendered "
  "as DNF on the dashboard. Retirements feed the Safety Car and RED-flag "
  "(mass-incident) Race Control triggers."),
 ("ev07_lap-completed", "lap-completed",
  "The car crosses the start/finish line and its lap counter increments.",
  "lapTime is computed (now - lap start) and a lap-completed event (details {lap, "
  "lapTime}) is published, persisted and shown on the dashboard. The leader's lap "
  "count is what drives the CHECKERED-flag trigger."),
 ("ev08_sector-completed", "sector-completed",
  "The car crosses one of the three Monza sector boundaries (S1 0-0.33, S2 "
  "0.33-0.66, S3 0.66-1.0).",
  "sectorTime is computed and a sector-completed event (details {sector, "
  "sectorTime, sectorName}) is published, persisted and used by the dashboard for "
  "live sector timing."),
 ("ev09_overtake", "overtake (reserved)",
  "A position swap between two cars on track. NOTE: 'overtake' is a DEFINED event "
  "type in event.schema.json and in the Mongo validator, but it is NOT currently "
  "emitted by any producer.",
  "Today, relative positions are recomputed by the Node-RED classification "
  "function (cars sorted by progress = lap + trackPos) and published, retained, "
  "on .../race-control/classification; the dashboard leaderboard re-renders and "
  "position swaps are shown implicitly. Comparing two consecutive standings is "
  "the natural extension point for emitting an explicit overtake event."),
]

FLAG = [
 ("ev10_GREEN", "GREEN",
  "Racing conditions. Set at the start of the race (reason race-start) and "
  "whenever Race Control clears a caution: the minimum flag duration has elapsed "
  "AND no car is in FAULT (hazard removed).",
  "changeFlag() validates the transition and the flag is published, retained, on "
  ".../race-control/flags (QoS 1). Node-RED caches it, appends it to the "
  "race_control history and broadcasts it. A GREEN with reason race-start also "
  "drives the cars INIT -> RUNNING."),
 ("ev11_YELLOW", "YELLOW",
  "Local caution for a single car stopped on track (faultCount == 1). It is a "
  "fallback that only fires when the VSC threshold is configured above 1; it is "
  "the only flag that carries a specific sector.",
  "Published retained with the affected sector; the dashboard highlights that "
  "track sector. It auto-clears to GREEN after the minimum YELLOW duration (10 s) "
  "once no car is in FAULT."),
 ("ev12_RED", "RED",
  "Serious multiple incident: count(FAULT + RETIRED) >= 3 (mass-incident "
  "threshold). The session is suspended.",
  "Published retained; the race is suspended and requires a MANUAL GREEN to "
  "restart (RED is not auto-cleared). While RED is active, Node-RED pauses the "
  "stationary-running detector so that legitimately stopped cars are not flagged "
  "as retired."),
 ("ev13_SC", "SC (Safety Car)",
  "Physical Safety Car deployed after a retirement that leaves debris. Each new "
  "retirement has a configurable probability (retirementScProbability, default 0, "
  "i.e. off unless enabled) of triggering it.",
  "Published retained; the dashboard shows the SC banner and cars are expected to "
  "slow down. It auto-clears to GREEN after the minimum SC duration (30 s) if no "
  "car is in FAULT."),
 ("ev14_VSC", "VSC (Virtual Safety Car)",
  "Virtual Safety Car for a hazard on track: faultCount >= multiFaultVscThreshold "
  "(default threshold = 1). This is the primary single-fault caution in the "
  "default configuration.",
  "Published retained; the dashboard shows the VSC banner. It auto-clears to "
  "GREEN after the minimum VSC duration (20 s) if no car is in FAULT."),
 ("ev15_CHECKERED", "CHECKERED",
  "End of race: the leader completes totalLaps (leaderLap >= totalLaps). It is "
  "the highest-priority trigger and an absorbing flag state.",
  "Published retained; no further flag transitions are allowed. Cars still in "
  "RUNNING/PIT transition to FINISHED (race-end) and the dashboard freezes the "
  "final classification."),
]

DERIVED = [
 ("ev16_tire-overheat", "tire-overheat (derived -> fault)",
  "Defensive cross-check computed in Node-RED (SF-Detect Telemetry Events) "
  "directly from the telemetry stream: the hottest of the four tyres reaches "
  ">= 170 C.",
  "After a 10 s per-car deduplication window, a synthetic car event of type "
  "'fault' is re-published on the events topic with details {reason: "
  "'tire-overheat', corner, tireTemp, threshold:170, source: "
  "'node-red-event-detection'} and meta.derived = true. It is persisted and shown "
  "on the dashboard like a native fault."),
 ("ev17_low-fuel", "low-fuel (derived -> state-change)",
  "Telemetry-side detection of a fuel emergency: fuel <= 5 kg while the car is "
  "not already in PIT, RETIRED or FINISHED.",
  "After a 30 s dedup window, a synthetic 'state-change' event (details {from: "
  "current state, to: 'PIT', reason: 'low-fuel'}) is published, persisted and "
  "broadcast, signalling that the car should pit."),
 ("ev18_stationary-running", "stationary-running (derived -> retirement)",
  "A car reported RUNNING but effectively stopped: speed <= 1 km/h held for "
  ">= 5 s. The check is suppressed while the race is RED (cars are legitimately "
  "stopped during a suspension).",
  "After a 10 s dedup window, a synthetic 'retirement' event (details {reason: "
  "'stationary-running', durationMs, speed, threshold:1}) is published, persisted "
  "and shown as a likely DNF on the dashboard."),
 ("ev19_track-position-jump", "track-position-jump (derived -> fault)",
  "Data-integrity / anti-cheat check on track progress: delta = (lap + trackPos) "
  "- previous is greater than 0.20 (forward teleport) or less than -0.05 "
  "(backstep), excluding the legitimate lap-0 race restart.",
  "After a 10 s dedup window, a synthetic 'fault' event (details {reason: "
  "'track-position-jump', delta, previousLap, currentLap, ...}) is published, "
  "persisted and surfaced on the dashboard for diagnostics."),
]

SUMMARY = [
 # (n, name, category, origin, trigger, emitted)
 ("1", "state-change", "Car lifecycle", "Simulator", "any accepted FSM transition", "state-change"),
 ("2", "pit-entry", "Car lifecycle", "Simulator", "tyre wear >=0.60 in pit window", "pit-entry"),
 ("3", "pit-stop", "Car lifecycle", "Simulator", "pit-box crossing (tyre change)", "pit-stop"),
 ("4", "pit-exit", "Car lifecycle", "Simulator", "pit timer 2.0-3.5 s elapsed", "pit-exit"),
 ("5", "fault", "Car lifecycle", "Simulator", "overheat / engine-failure / internal-error", "fault"),
 ("6", "retirement", "Car lifecycle", "Simulator", "FAULT 5 s / manual-retire", "retirement"),
 ("7", "lap-completed", "Car lifecycle", "Simulator", "start/finish crossing", "lap-completed"),
 ("8", "sector-completed", "Car lifecycle", "Simulator", "sector boundary crossing", "sector-completed"),
 ("9", "overtake", "Car lifecycle", "Reserved (schema)", "position swap (not emitted)", "overtake*"),
 ("10", "GREEN", "Race Control", "Auto/Manual", "race-start / caution cleared", "flag GREEN"),
 ("11", "YELLOW", "Race Control", "Auto/Manual", "single car FAULT (count==1)", "flag YELLOW"),
 ("12", "RED", "Race Control", "Auto/Manual", "FAULT+RETIRED >= 3", "flag RED"),
 ("13", "SC", "Race Control", "Auto/Manual", "retirement w/ probability", "flag SC"),
 ("14", "VSC", "Race Control", "Auto/Manual", "faultCount >= 1", "flag VSC"),
 ("15", "CHECKERED", "Race Control", "Auto/Manual", "leaderLap >= totalLaps", "flag CHECKERED"),
 ("16", "tire-overheat", "Derived (Node-RED)", "Node-RED", "hottest tyre >= 170 C", "fault"),
 ("17", "low-fuel", "Derived (Node-RED)", "Node-RED", "fuel <= 5 kg (not in pit)", "state-change"),
 ("18", "stationary-running", "Derived (Node-RED)", "Node-RED", "speed <=1 km/h for >=5 s", "retirement"),
 ("19", "track-position-jump", "Derived (Node-RED)", "Node-RED", "progress delta >0.20 / <-0.05", "fault"),
]


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)


def add_figure(doc, img_name):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(os.path.join(HERE, img_name + ".png"), width=IMG_W)


def add_event(doc, idx, img, name, why, handling, fig_no):
    doc.add_heading("%s. %s" % (idx, name), level=3)
    p = doc.add_paragraph()
    p.add_run("Why it happens. ").bold = True
    p.add_run(why)
    p = doc.add_paragraph()
    p.add_run("Handling. ").bold = True
    p.add_run(handling)
    add_figure(doc, img)
    add_caption(doc, "Figure A.%s - %s sequence (why + handling)" % (fig_no, name))


def main():
    doc = docx.Document(REPORT)

    doc.add_page_break()
    doc.add_heading("Appendix A - Event Catalogue", level=1)
    doc.add_paragraph(
        "This appendix documents every discrete event the system can produce: 9 "
        "car-lifecycle events emitted by the simulator (schemas/event.schema.json), "
        "6 Race Control flags (schemas/flag.schema.json) and 4 derived telemetry "
        "events computed defensively in Node-RED - 19 in total. For each event we "
        "give why it happens, how it is handled across the pipeline and a UML "
        "sequence diagram. Car events travel on "
        "f1/simulation/{raceId}/teams/{teamId}/cars/{carId}/events (QoS 1); flags "
        "travel, retained, on f1/simulation/{raceId}/race-control/flags."
    )

    # Summary table
    doc.add_heading("A.0 Summary", level=2)
    cols = ["#", "Event", "Category", "Origin", "Trigger (why)", "Type/flag emitted"]
    tbl = doc.add_table(rows=1, cols=len(cols))
    tbl.style = "Normal Table"
    set_table_borders(tbl)
    for i, c in enumerate(cols):
        run = tbl.rows[0].cells[i].paragraphs[0].add_run(c)
        run.bold = True
    for row in SUMMARY:
        cells = tbl.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = val
    doc.add_paragraph(
        "* overtake is defined in the schema and Mongo validator but is not "
        "currently emitted; see A.1.9.", style="Normal")

    fig = 1
    doc.add_heading("A.1 Car Lifecycle Events (Simulator)", level=2)
    doc.add_paragraph(
        "Emitted by the per-car finite state machine (simulator/src/car/). Each "
        "transition is decided in conditions.js, applied by fsm.js and turned into "
        "MQTT events by Car.tick().")
    for i, (img, name, why, handling) in enumerate(CAR, start=1):
        add_event(doc, "A.1.%d" % i, img, name, why, handling, fig); fig += 1

    doc.add_heading("A.2 Race Control Flags", level=2)
    doc.add_paragraph(
        "Global session events. The automatic triggers (triggers.js) are evaluated "
        "every orchestrator tick in priority order CHECKERED > RED > SC > VSC > "
        "clearance; flags can also be set manually from the dashboard through the "
        "AES-256-GCM encrypted command channel (.../secure/race-control/flags). "
        "flag-state.js enforces the allowed transition matrix.")
    for i, (img, name, why, handling) in enumerate(FLAG, start=1):
        add_event(doc, "A.2.%d" % i, img, name, why, handling, fig); fig += 1

    doc.add_heading("A.3 Derived Telemetry Events (Node-RED)", level=2)
    doc.add_paragraph(
        "Computed in the SF-Detect Telemetry Events subflow directly from the "
        "telemetry stream, independently of the simulator FSM. They reuse the car "
        "event schema (type fault / state-change / retirement) but are tagged "
        "meta.derived = true and carry source = 'node-red-event-detection'; each "
        "reason has its own deduplication window.")
    for i, (img, name, why, handling) in enumerate(DERIVED, start=1):
        add_event(doc, "A.3.%d" % i, img, name, why, handling, fig); fig += 1

    out = REPORT
    try:
        doc.save(out)
        print("SAVED:", out)
    except PermissionError:
        out = REPORT.replace(".docx", "_with_events.docx")
        doc.save(out)
        print("LOCKED original (Word open). SAVED COPY:", out)


if __name__ == "__main__":
    main()
